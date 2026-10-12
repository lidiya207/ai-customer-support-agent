"""Extract plain text from uploaded knowledge files."""

from __future__ import annotations

import io
from pathlib import Path

from fastapi import HTTPException


def extract_text_from_upload(filename: str, raw: bytes) -> tuple[str, str]:
    """
    Return (plain_text, doc_type_name) where doc_type_name is TXT | PDF | DOCX.
    """
    lower = filename.lower()
    suffix = Path(lower).suffix

    if suffix in {".txt", ".md"}:
        try:
            return raw.decode("utf-8"), "TXT"
        except UnicodeDecodeError as exc:
            raise HTTPException(
                status_code=400, detail="Text files must be UTF-8 encoded."
            ) from exc

    if suffix == ".pdf":
        return _extract_pdf(raw), "PDF"

    if suffix == ".docx":
        return _extract_docx(raw), "DOCX"

    raise HTTPException(
        status_code=400,
        detail="Only .txt, .md, .pdf, and .docx uploads are supported.",
    )


def _extract_pdf(raw: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="PDF support is not installed. Install pypdf.",
        ) from exc

    try:
        reader = PdfReader(io.BytesIO(raw))
        parts: list[str] = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text.strip())
        content = "\n\n".join(parts).strip()
    except Exception as exc:  # noqa: BLE001 - surface as upload error
        raise HTTPException(
            status_code=400, detail=f"Could not read PDF: {exc}"
        ) from exc

    if not content:
        raise HTTPException(
            status_code=400,
            detail="No extractable text found in this PDF (it may be scanned/image-only).",
        )
    return content


def _extract_docx(raw: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="DOCX support is not installed. Install python-docx.",
        ) from exc

    try:
        document = Document(io.BytesIO(raw))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        # Include simple table cell text
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        content = "\n\n".join(parts).strip()
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(
            status_code=400, detail=f"Could not read DOCX: {exc}"
        ) from exc

    if not content:
        raise HTTPException(
            status_code=400, detail="No extractable text found in this DOCX."
        )
    return content
